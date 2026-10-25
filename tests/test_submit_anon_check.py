"""Tests for tools/submit/anon_check.py.

Each case is an anonymity leak a reviewer sees in the file properties while the author
sees nothing in the Word window. Needs python-docx (skipped when missing); the image cases
need Pillow and exiftool; the PDF case needs PyMuPDF."""
import json
import shutil

import pytest

from conftest import run_tool

docx = pytest.importorskip("docx")
from docx import Document  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls  # noqa: E402

HAS_EXIF = shutil.which("exiftool") is not None
NAMES = ["--names", "Jane Doe,Doe, J.,陳小明"]


def make_jpeg(path, artist=None, gps=False):
    Image = pytest.importorskip("PIL.Image")
    im = Image.new("RGB", (40, 40), (200, 30, 30))
    ex = Image.Exif()
    if artist:
        ex[0x013B] = artist
    if gps:
        g = ex.get_ifd(0x8825)
        g[1], g[2], g[3], g[4] = "N", (25.0, 2.0, 30.0), "E", (121.0, 30.0, 0.0)
    im.save(str(path), "JPEG", exif=ex)


def clean_doc(path, text="This is a clean body paragraph. 這是一段乾淨的內文。"):
    d = Document()
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    d.add_paragraph(text)
    d.save(str(path))
    return d


def cli(*args):
    return run_tool("submit/anon_check.py", *args)


def test_clean_docx_passes(tmp_path):
    clean_doc(tmp_path / "a.docx")
    r = cli(tmp_path / "a.docx", *NAMES, "--affil", "Example University,EU")
    assert r.returncode == 0, r.stdout + r.stderr
    assert "🔴 0" in r.stdout.splitlines()[-1]
    assert "Actually checked:" in r.stdout
    assert "empty run" not in r.stdout


def test_core_creator_name(tmp_path):
    d = clean_doc(tmp_path / "a.docx")
    d.core_properties.author = "Jane Doe"
    d.save(str(tmp_path / "a.docx"))
    r = cli(tmp_path / "a.docx", *NAMES)
    assert r.returncode == 1, r.stdout
    assert "core.xml creator" in r.stdout
    assert 'matched name "Jane Doe"' in r.stdout


def test_core_unknown_creator_is_orange_only(tmp_path):
    d = clean_doc(tmp_path / "a.docx")
    d.core_properties.author = "user123"
    d.save(str(tmp_path / "a.docx"))
    r = cli(tmp_path / "a.docx", *NAMES)
    assert r.returncode == 0, r.stdout
    assert "🟠 [docProps/core.xml creator]" in r.stdout


def test_body_affiliation(tmp_path):
    clean_doc(tmp_path / "a.docx", "This study was carried out at Example University, Department of Design.")
    r = cli(tmp_path / "a.docx", *NAMES, "--affil", "Example University")
    assert r.returncode == 1
    assert 'matched affiliation "Example University"' in r.stdout
    assert "【Example University】" in r.stdout


def test_name_reverse_order_and_case(tmp_path):
    clean_doc(tmp_path / "a.docx", "Contact: DOE Jane for details.")
    r = cli(tmp_path / "a.docx", *NAMES)
    assert r.returncode == 1, r.stdout


def test_cjk_name_in_body(tmp_path):
    clean_doc(tmp_path / "a.docx", "感謝陳小明協助訪談。")
    assert cli(tmp_path / "a.docx", *NAMES).returncode == 1


def test_hyperlink_account(tmp_path):
    d = clean_doc(tmp_path / "a.docx")
    rid = d.part.relate_to("https://github.com/janedoe/proj", RT.HYPERLINK, is_external=True)
    d.paragraphs[0]._p.append(parse_xml(
        '<w:hyperlink %s r:id="%s"><w:r><w:t>code</w:t></w:r></w:hyperlink>' % (nsdecls("w", "r"), rid)))
    d.save(str(tmp_path / "a.docx"))
    r = cli(tmp_path / "a.docx", *NAMES, "--extra", "janedoe,example.org")
    assert r.returncode == 1
    assert 'matched extra string "janedoe"' in r.stdout
    assert ".rels" in r.stdout


def test_leftover_comment(tmp_path):
    d = clean_doc(tmp_path / "a.docx")
    if not hasattr(d, "add_comment"):
        pytest.skip("this python-docx has no add_comment")
    d.add_comment(d.paragraphs[0].runs, text="Fix this", author="Some Reviewer", initials="X")
    d.save(str(tmp_path / "a.docx"))
    r = cli(tmp_path / "a.docx", *NAMES)
    assert r.returncode == 1, r.stdout
    assert "1 leftover comments" in r.stdout


def test_leftover_tracked_change(tmp_path):
    d = clean_doc(tmp_path / "a.docx")
    d.paragraphs[0]._p.append(parse_xml(
        '<w:ins %s w:id="9" w:author="Someone" w:date="2026-01-01T00:00:00Z"><w:r><w:t>added</w:t></w:r></w:ins>' % nsdecls("w")))
    d.save(str(tmp_path / "a.docx"))
    r = cli(tmp_path / "a.docx", *NAMES)
    assert r.returncode == 1, r.stdout
    assert "1 leftover tracked changes" in r.stdout


def test_self_disclosure_is_orange_only(tmp_path):
    clean_doc(tmp_path / "a.docx", "As shown earlier, our previous work holds. 筆者先前的研究亦然。")
    r = cli(tmp_path / "a.docx", *NAMES)
    assert r.returncode == 0, r.stdout
    assert "English self-disclosure" in r.stdout
    assert "Chinese self-disclosure" in r.stdout
    assert "🟠 2" in r.stdout.splitlines()[-1]


@pytest.mark.skipif(not HAS_EXIF, reason="needs exiftool")
def test_embedded_image_exif(tmp_path):
    make_jpeg(tmp_path / "pic.jpg", artist="Jane Doe", gps=True)
    d = clean_doc(tmp_path / "a.docx")
    d.add_picture(str(tmp_path / "pic.jpg"))
    d.save(str(tmp_path / "a.docx"))
    r = cli(tmp_path / "a.docx", *NAMES)
    assert r.returncode == 1, r.stdout
    assert "word/media/" in r.stdout
    assert "Artist" in r.stdout
    assert "GPS" in r.stdout


@pytest.mark.skipif(not HAS_EXIF, reason="needs exiftool")
def test_image_file_and_folder(tmp_path):
    (tmp_path / "sub").mkdir()
    make_jpeg(tmp_path / "sub" / "x.jpg", artist="Someone Else")
    r = cli(tmp_path / "sub", *NAMES)
    assert r.returncode == 1, r.stdout       # any value in an Artist field is a red finding
    assert "author / copyright field" in r.stdout


def test_exiftool_missing_degrades_to_orange(tmp_path, monkeypatch):
    # PATH without exiftool: the EXIF area must be reported as not checked, not silently passed
    pytest.importorskip("PIL.Image")
    make_jpeg(tmp_path / "pic.jpg", artist="Jane Doe")
    d = clean_doc(tmp_path / "a.docx")
    d.add_picture(str(tmp_path / "pic.jpg"))
    d.save(str(tmp_path / "a.docx"))
    from conftest import load_tool
    ac = load_tool("submit/anon_check.py")
    monkeypatch.setattr(ac, "exiftool_path", lambda: None)
    rep = ac.check_docx(str(tmp_path / "a.docx"), ac.build_terms("Jane Doe", "", ""))
    rep.finish()
    assert any("exiftool not found" in f[2] for f in rep.findings)


def test_pdf_author(tmp_path):
    fitz = pytest.importorskip("pymupdf")
    doc = fitz.open()
    doc.new_page()
    doc.set_metadata({"author": "Jane Doe", "title": "Anonymous manuscript"})
    doc.save(str(tmp_path / "a.pdf"))
    r = cli(tmp_path / "a.pdf", *NAMES)
    assert r.returncode == 1, r.stdout
    assert "Info Author" in r.stdout
    assert "empty run" in r.stdout          # a blank page yields no text: say so loudly


def test_unreadable_exit2(tmp_path):
    (tmp_path / "bad.docx").write_bytes(b"not a zip")
    r = cli(tmp_path / "bad.docx", *NAMES)
    assert r.returncode == 2
    assert "unreadable" in r.stdout
    assert cli(tmp_path / "missing.pdf", *NAMES).returncode == 2


def test_json_output(tmp_path):
    clean_doc(tmp_path / "a.docx")
    r = cli(tmp_path / "a.docx", *NAMES, "--json")
    j = json.loads(r.stdout)
    assert j["red"] == 0
    assert j["files"][0]["total_checked"] > 0
