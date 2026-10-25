"""Tests for tools/collab/docx_diff.py and tools/collab/coding_pack.py.

docx files are generated on the fly with python-docx; coding data is fake and lives in tmp_path.
Chinese appears only in fixtures, where the tools must handle Chinese text.
"""
import csv
import json
import os
import re
import shutil
import subprocess

import pytest

from conftest import load_tool, run_tool

DIFF = "collab/docx_diff.py"
PACK = "collab/coding_pack.py"

# ---------------------------------------------------------------- docx_diff

BASE = [("h", "一、研究背景"), ("p", "本計畫探討城市中的非人類生物。"), ("p", "我們採用參與式設計方法進行田野調查。"),
        ("h", "二、研究方法"), ("p", "第一階段蒐集資料，第二階段進行分析。"), ("p", "最後提出設計建議與討論。")]


@pytest.fixture
def docx_lib():
    pytest.importorskip("lxml")
    return pytest.importorskip("docx")


def make(docx_lib, path, items, table=None, comment=None, tracked=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = docx_lib.Document()
    for kind, text in items:
        if kind == "h":
            d.add_heading(text, level=1)
        else:
            d.add_paragraph(text)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
    if comment:
        if not hasattr(d, "add_comment"):
            pytest.skip("this python-docx version has no Document.add_comment")
        para = d.paragraphs[1]
        d.add_comment(para.runs, text=comment, author="Reviewer One")
    if tracked:
        p = d.paragraphs[1]._p
        ins = OxmlElement("w:ins")
        ins.set(qn("w:id"), "901")
        ins.set(qn("w:author"), "Reviewer One")
        ins.set(qn("w:date"), "2026-10-01T00:00:00Z")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "(tracked insertion)"
        r.append(t)
        ins.append(r)
        p.append(ins)
        dl = OxmlElement("w:del")
        dl.set(qn("w:id"), "902")
        dl.set(qn("w:author"), "Reviewer One")
        dl.set(qn("w:date"), "2026-10-01T00:00:00Z")
        r2 = OxmlElement("w:r")
        dt = OxmlElement("w:delText")
        dt.text = "old"
        r2.append(dt)
        dl.append(r2)
        p.append(dl)
    d.save(str(path))
    return str(path)


def diff_pair(docx_lib, tmp_path, new_items, old_table=None, **kw):
    old = make(docx_lib, tmp_path / "old.docx", BASE, table=old_table)
    new = make(docx_lib, tmp_path / "new.docx", new_items, **kw)
    js = tmp_path / "c.json"
    r = run_tool(DIFF, old, new, "--json", js)
    data = json.loads(js.read_text(encoding="utf-8")) if js.exists() else []
    return r, data


def sub(items, idx, text):
    out = list(items)
    out[idx] = (out[idx][0], text)
    return out


def test_identical_exit0(docx_lib, tmp_path):
    r, data = diff_pair(docx_lib, tmp_path, BASE)
    assert r.returncode == 0 and data == []
    assert "identical content" in r.stdout


def test_word_changed(docx_lib, tmp_path):
    r, data = diff_pair(docx_lib, tmp_path, sub(BASE, 1, "本計畫探討城市中的非人類動物。"))
    assert r.returncode == 1 and len(data) == 1
    c = data[0]
    assert c["type"] == "modified" and c["heading"] == "一、研究背景"
    assert "[-生-]" in c["diff"] and "{+動+}" in c["diff"] and not c["punct_only"]


def test_paragraph_deleted(docx_lib, tmp_path):
    r, data = diff_pair(docx_lib, tmp_path, BASE[:2] + BASE[3:])
    assert [c["type"] for c in data] == ["deleted"]
    assert "參與式設計" in data[0]["old"] and data[0]["old_index"] == 3


def test_paragraph_added(docx_lib, tmp_path):
    items = BASE[:3] + [("p", "另外加入焦點團體訪談作為補充資料來源。")] + BASE[3:]
    r, data = diff_pair(docx_lib, tmp_path, items)
    assert [c["type"] for c in data] == ["added"]
    assert data[0]["heading"] == "一、研究背景" and "焦點團體" in data[0]["new"]


def test_paragraph_moved(docx_lib, tmp_path):
    items = [BASE[0], BASE[2], BASE[1]] + BASE[3:]
    r, data = diff_pair(docx_lib, tmp_path, items)
    types = [c["type"] for c in data]
    assert "moved" in types and "added" not in types and "deleted" not in types


def test_table_cell_changed(docx_lib, tmp_path):
    tb = [["項目", "金額"], ["人事費", "100"]]
    tb2 = [["項目", "金額"], ["人事費", "120"]]
    r, data = diff_pair(docx_lib, tmp_path, BASE, old_table=tb, table=tb2)
    assert [c["type"] for c in data] == ["table cell modified"]
    assert data[0]["old"] == "100" and data[0]["new"] == "120" and "table 1" in data[0]["note"]


def test_punctuation_only(docx_lib, tmp_path):
    r, data = diff_pair(docx_lib, tmp_path, sub(BASE, 4, "第一階段蒐集資料；第二階段進行分析。"))
    assert len(data) == 1 and data[0]["type"] == "modified" and data[0]["punct_only"] is True
    assert "folded" in r.stdout


def test_heading_changed(docx_lib, tmp_path):
    r, data = diff_pair(docx_lib, tmp_path, sub(BASE, 3, "二、研究方法與步驟"))
    assert len(data) == 1 and data[0]["type"] == "modified"
    assert data[0]["heading"] == "二、研究方法與步驟"


def test_comment_reported(docx_lib, tmp_path):
    r, data = diff_pair(docx_lib, tmp_path, BASE, comment="Please add the species names here")
    assert r.returncode == 0          # a comment is not a content difference
    assert "Please add the species names here" in r.stdout and "Reviewer One" in r.stdout
    assert "本計畫探討" in r.stdout   # the anchored text is shown


def test_tracked_changes_detected(docx_lib, tmp_path):
    r, data = diff_pair(docx_lib, tmp_path, BASE, tracked=True)
    assert "**has** tracked changes" in r.stdout and "Reviewer One" in r.stdout
    assert "1 insertions, 1 deletions" in r.stdout


def test_no_tracking_stated(docx_lib, tmp_path):
    r, _ = diff_pair(docx_lib, tmp_path, BASE)
    assert "the new file has none" in r.stdout


def test_out_report_file(docx_lib, tmp_path):
    old = make(docx_lib, tmp_path / "old.docx", BASE)
    new = make(docx_lib, tmp_path / "new.docx", sub(BASE, 1, "本計畫探討城市中的非人類動物。"))
    rep = tmp_path / "report.md"
    r = run_tool(DIFF, old, new, "--out", rep)
    assert r.returncode == 1
    assert "## Change list" in rep.read_text(encoding="utf-8")


def test_bad_file_exit2(docx_lib, tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_text("not a zip")
    ok = make(docx_lib, tmp_path / "ok.docx", BASE)
    r = run_tool(DIFF, bad, ok)
    assert r.returncode == 2


# ---------------------------------------------------------------- coding_pack

CODEBOOK = """# Codebook

## General rules
1. Several codes may apply.
2. Use "other" only when none apply.

## Question 1: What is the most important thing you learned?

| Code | Definition | Boundary (what does not count) |
|---|---|---|
| L_teamwork | mentions "working together", division of labour | individual work only |
| L_sound (new) | sound｜speech (full-width bar), audio \\| across columns | not music |
| L_other | none of the above | |

## Question 2: Any other thoughts?

| Code | Definition | Boundary |
|---|---|---|
| T_fun | enjoyable | |
| T_hard | too hard | |
"""

CODEBOOK_ZH = """# 編碼簿

## 總則
1. 可複選。

## 題一：您學到最重要的一件事？

| 碼 | 定義 | 邊界（何者不算） |
|---|---|---|
| L_團隊協作 | 提到合作 | 只提個人 |
| L_聲音知識（新增） | 提到聲音 | 不含音樂 |
| L_其他 | 以上皆非 | |
"""

SECRET = "unique-answer-text-XYZ"
Q1_CODES = ["L_teamwork", "L_sound", "L_other"]
Q2_CODES = ["T_fun", "T_hard"]


@pytest.fixture
def cp():
    pytest.importorskip("numpy")
    return load_tool(PACK)


def write_csv(path, header, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


@pytest.fixture
def pack(tmp_path):
    pytest.importorskip("numpy")
    cb = tmp_path / "codebook.md"
    cb.write_text(CODEBOOK, encoding="utf-8")
    c1 = tmp_path / "coderB_Second_BLANK_q1.csv"
    c2 = tmp_path / "coderB_Second_BLANK_q2.csv"
    write_csv(c1, ["pid", "response"] + Q1_CODES, [[f"P{i}", f"{SECRET}{i}", "", "", ""] for i in range(4)])
    write_csv(c2, ["pid", "response"] + Q2_CODES, [[f"P{i}", f"second answer {i}", "", ""] for i in range(4)])
    return {"dir": tmp_path, "cb": str(cb), "c1": str(c1), "c2": str(c2)}


def build(pack, out=None, extra=()):
    out = out or str(pack["dir"] / "out" / "pack.html")
    return run_tool(PACK, "build", "--codebook", pack["cb"], "--csv", pack["c1"], "--csv", pack["c2"],
                    "--coder", "Second", "--out", out, "--allow-tracked", *extra), out


def pack_hash_of(html_path):
    with open(html_path, encoding="utf-8") as f:
        return re.search(r'"hash": "([0-9a-f]{64})"', f.read()).group(1)


def make_export(pack, h, values=None, coder="Second", drop=None):
    """values: {(file, pid, code): value}; by default the first non-other code is ticked."""
    rows = [["pack_hash", "coder", "question", "pid", "code", "value"]]
    spec = [(os.path.basename(pack["c1"]), Q1_CODES), (os.path.basename(pack["c2"]), Q2_CODES)]
    for f, codes in spec:
        for i in range(4):
            for j, c in enumerate(codes):
                v = 1 if j == 0 else 0
                v = (values or {}).get((f, f"P{i}", c), v)
                if drop and drop == (f, f"P{i}", c):
                    continue
                rows.append([h, coder, f, f"P{i}", c, str(v)])
    p = pack["dir"] / "export.csv"
    with open(p, "w", encoding="utf-8-sig", newline="") as fh:
        csv.writer(fh).writerows(rows)
    return str(p)


def _pair(tmp_path, a_vals, b_vals, code="C", name="q"):
    ids = [f"P{i}" for i in range(len(a_vals))]
    pa, pb = tmp_path / f"A_DONE_{name}.csv", tmp_path / f"B_DONE_{name}.csv"
    write_csv(pa, ["pid", "response", code], [[i, f"{SECRET}{i}", v] for i, v in zip(ids, a_vals)])
    write_csv(pb, ["pid", "response", code], [[i, f"{SECRET}{i}", v] for i, v in zip(ids, b_vals)])
    return str(pa), str(pb)


def hand_example():
    # n11=3, n10=1, n01=1, n00=5: po=.8 pe=.52 -> kappa=0.58333, PABAK=0.6
    a = [1, 1, 1, 1, 0, 0, 0, 0, 0, 0]
    b = [1, 1, 1, 0, 1, 0, 0, 0, 0, 0]
    return a, b


def read_summary(d):
    with open(os.path.join(d, "kappa_summary.csv"), encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def test_parse_codebook_suffix_rules_and_symbols(cp):
    rules, qs = cp.parse_codebook(CODEBOOK)
    assert "Several codes may apply" in rules and "other" in rules
    assert [q["label"] for q in qs] == ["Question 1", "Question 2"]
    assert [c["name"] for c in qs[0]["codes"]] == ["L_teamwork", "L_sound", "L_other"]
    assert "｜" in qs[0]["codes"][1]["definition"]
    assert "|" in qs[0]["codes"][1]["definition"]  # the \| escape is restored
    assert qs[0]["codes"][0]["boundary"] == "individual work only"
    assert cp.display_name("L_teamwork") == "teamwork"


def test_parse_codebook_chinese(cp):
    rules, qs = cp.parse_codebook(CODEBOOK_ZH)
    assert "可複選" in rules
    assert qs[0]["label"] == "題一" and qs[0]["text"] == "您學到最重要的一件事？"
    assert [c["name"] for c in qs[0]["codes"]] == ["L_團隊協作", "L_聲音知識", "L_其他"]
    assert cp.is_other("L_其他") and cp.is_other("L_other") and not cp.is_other("L_otherwise")


def test_plain_md_strips_inline_marks(cp):
    # the page does not render markdown, so inline marks must be removed from the rules and definitions
    assert cp.plain_md("**bold** and `other` and [link](http://x)") == "bold and other and link"
    assert cp.plain_md("2. *none apply* then") == "2. none apply then"
    assert cp.plain_md("a*b*c 2*3=6") == "a*b*c 2*3=6"
    assert cp.plain_md("* list item") == "* list item"
    rules, qs = cp.parse_codebook(
        "## General rules\n1. **Several** codes.\n\n## Question 1: x\n| Code | Definition | Boundary |\n|---|---|---|\n"
        "| `L_a` | has `code` | **only** individual |\n")
    assert rules == "1. Several codes."
    assert qs[0]["codes"][0]["definition"] == "has code"
    assert qs[0]["codes"][0]["boundary"] == "only individual"


def test_import_into_missing_dir_exit2(tmp_path):
    pytest.importorskip("numpy")
    r = run_tool(PACK, "import", tmp_path / "x.csv", "--into", tmp_path / "nope")
    assert r.returncode == 2


def test_codebook_without_questions_fails(cp):
    with pytest.raises(cp.Fail):
        cp.parse_codebook("# x\n\n## General rules\nrules only\n")


def test_build_ok_selfcontained_and_js_syntax(pack):
    r, out = build(pack)
    assert r.returncode == 0, r.stderr
    page = open(out, encoding="utf-8").read()
    assert "http://" not in page and "https://" not in page and "src=" not in page
    assert SECRET in page  # the response text is inside the HTML
    assert "contains participants' answers" in r.stderr
    assert pack_hash_of(out)[:12] in page
    assert '<html lang="en">' in page and "Second-coder coding" in page
    if not shutil.which("node"):
        pytest.skip("node is not installed; embedded JavaScript syntax not checked")
    js = re.search(r"<script>\n(.*?)</script>", page, re.S).group(1)
    jp = pack["dir"] / "x.js"
    jp.write_text(js, encoding="utf-8")
    chk = subprocess.run(["node", "--check", str(jp)], capture_output=True, text=True)
    assert chk.returncode == 0, chk.stderr


def test_build_ui_zh_tw(pack):
    r, out = build(pack, extra=("--ui", "zh-TW"))
    assert r.returncode == 0, r.stderr
    page = open(out, encoding="utf-8").read()
    assert '<html lang="zh-Hant">' in page and "第二編碼者判讀" in page and "匯出 CSV" in page
    assert "Export CSV" not in page
    assert pack_hash_of(out) == pack_hash_of(build(pack, str(pack["dir"] / "en.html"))[1])


def test_build_ui_unknown_language_rejected(pack):
    r, _ = build(pack, extra=("--ui", "fr"))
    assert r.returncode == 2  # argparse rejects an unknown choice


def test_page_has_no_banned_dashes(pack):
    r, out = build(pack)
    page = open(out, encoding="utf-8").read()
    banned = "".join(chr(c) for c in (0x2014, 0x2015, 0x2E3A, 0x2E3B))
    assert not re.search("[" + banned + "]", page)


def test_build_code_mismatch_exit2(pack):
    write_csv(pack["c2"], ["pid", "response", "T_fun", "T_extra"], [["P0", "x", "", ""]])
    r, _ = build(pack)
    assert r.returncode == 2
    assert "T_hard" in r.stderr and "T_extra" in r.stderr


def test_build_csv_count_mismatch(pack):
    r = run_tool(PACK, "build", "--codebook", pack["cb"], "--csv", pack["c1"], "--coder", "x",
                 "--out", pack["dir"] / "p.html", "--allow-tracked")
    assert r.returncode == 2


def test_hash_stable_and_sensitive(pack):
    _, o1 = build(pack, str(pack["dir"] / "a.html"))
    _, o2 = build(pack, str(pack["dir"] / "b.html"))
    assert pack_hash_of(o1) == pack_hash_of(o2)
    write_csv(pack["c2"], ["pid", "response"] + Q2_CODES, [[f"P{i}", f"edited {i}", "", ""] for i in range(4)])
    _, o3 = build(pack, str(pack["dir"] / "c.html"))
    assert pack_hash_of(o3) != pack_hash_of(o1)


def test_privacy_gate_blocks_untracked_git_path(pack):
    if not shutil.which("git"):
        pytest.skip("git is not installed")
    repo = pack["dir"] / "repo"
    repo.mkdir()
    subprocess.run(["git", "init", "-q"], cwd=repo, check=True)
    out = str(repo / "sub" / "pack.html")
    args = ("build", "--codebook", pack["cb"], "--csv", pack["c1"], "--csv", pack["c2"],
            "--coder", "Second", "--out", out)
    r = run_tool(PACK, *args)
    assert r.returncode == 3
    assert not os.path.exists(out)
    (repo / ".gitignore").write_text("*.html\n", encoding="utf-8")
    r = run_tool(PACK, *args)
    assert r.returncode == 0, r.stderr
    assert os.path.exists(out)


def test_privacy_gate_allow_tracked(pack):
    if not shutil.which("git"):
        pytest.skip("git is not installed")
    repo = pack["dir"] / "repo2"
    repo.mkdir()
    subprocess.run(["git", "init", "-q"], cwd=repo, check=True)
    out = str(repo / "pack.html")
    r = run_tool(PACK, "build", "--codebook", pack["cb"], "--csv", pack["c1"], "--csv", pack["c2"],
                 "--coder", "Second", "--out", out, "--allow-tracked")
    assert r.returncode == 0
    assert "warning" in r.stderr


def test_import_ok_and_refuse_overwrite(cp, pack):
    _, out = build(pack)
    ex = make_export(pack, pack_hash_of(out))
    r = run_tool(PACK, "import", ex, "--into", pack["dir"])
    assert r.returncode == 0, r.stderr
    done = pack["dir"] / "coderB_Second_DONE_q1.csv"
    assert done.exists()
    header, rows = cp.read_csv(str(done))
    assert header == ["pid", "response"] + Q1_CODES
    assert rows[0][1].startswith(SECRET)  # response text kept exactly as in the BLANK file
    assert rows[0][2:] == ["1", "0", "0"]
    r2 = run_tool(PACK, "import", ex, "--into", pack["dir"])
    assert r2.returncode == 4
    r3 = run_tool(PACK, "import", ex, "--into", pack["dir"], "--force")
    assert r3.returncode == 0


def test_import_rejects_missing_value(pack):
    _, out = build(pack)
    f = os.path.basename(pack["c1"])
    ex = make_export(pack, pack_hash_of(out), drop=(f, "P2", "L_other"))
    r = run_tool(PACK, "import", ex, "--into", pack["dir"])
    assert r.returncode == 2 and "P2" in r.stderr and "missing code" in r.stderr
    assert not (pack["dir"] / "coderB_Second_DONE_q1.csv").exists()


def test_import_rejects_hash_mismatch(pack):
    _, out = build(pack)
    ex = make_export(pack, "0" * 64)
    r = run_tool(PACK, "import", ex, "--into", pack["dir"])
    assert r.returncode == 2 and "hash" in r.stderr


def test_import_rejects_hash_when_blank_modified(pack):
    _, out = build(pack)
    ex = make_export(pack, pack_hash_of(out))
    write_csv(pack["c1"], ["pid", "response"] + Q1_CODES, [[f"P{i}", "edited", "", "", ""] for i in range(4)])
    assert run_tool(PACK, "import", ex, "--into", pack["dir"]).returncode == 2


def test_import_rejects_non_binary_value(pack):
    _, out = build(pack)
    f = os.path.basename(pack["c1"])
    ex = make_export(pack, pack_hash_of(out), values={(f, "P0", "L_teamwork"): 2})
    r = run_tool(PACK, "import", ex, "--into", pack["dir"])
    assert r.returncode == 2 and "not 0/1" in r.stderr


def test_import_other_exclusivity(pack):
    _, out = build(pack)
    h = pack_hash_of(out)
    f = os.path.basename(pack["c1"])
    # "other" and another code both 1
    ex = make_export(pack, h, values={(f, "P0", "L_other"): 1})
    r = run_tool(PACK, "import", ex, "--into", pack["dir"])
    assert r.returncode == 2 and "mutually exclusive" in r.stderr
    # all 0 while "other" is 0
    ex = make_export(pack, h, values={(f, "P1", "L_teamwork"): 0})
    r = run_tool(PACK, "import", ex, "--into", pack["dir"])
    assert r.returncode == 2 and "L_other" in r.stderr
    # a correct "other only" row passes
    ex = make_export(pack, h, values={(f, "P1", "L_teamwork"): 0, (f, "P1", "L_other"): 1})
    assert run_tool(PACK, "import", ex, "--into", pack["dir"]).returncode == 0


def test_score_kappa_hand_example(tmp_path):
    pytest.importorskip("numpy")
    a, b = hand_example()
    pa, pb = _pair(tmp_path, a, b)
    od = tmp_path / "o"
    r = run_tool(PACK, "score", pa, pb, "--threshold", "0.5", "--low-prev", "0.10", "--boot", "300", "--out-dir", od)
    assert r.returncode == 0, r.stderr
    row = read_summary(od)[0]
    assert abs(float(row["kappa"]) - 0.583) < 0.001
    assert abs(float(row["pabak"]) - 0.6) < 0.001
    assert row["pct_agree"] == "80.0"
    assert row["flag"] == "meets threshold" and row["a_checked"] == "4" and row["b_checked"] == "4"
    assert float(row["ci_low"]) < 0.583 < float(row["ci_high"])
    assert SECRET not in r.stdout and SECRET not in r.stderr  # never print response text
    assert "before seeing the data" in r.stdout
    with open(od / "disagreements_A_DONE_q.csv", encoding="utf-8-sig") as fh:
        dis = list(csv.DictReader(fh))
    assert sorted(d["pid"] for d in dis) == ["P3", "P4"]
    assert set(dis[0]) == {"pid", "code", "A", "B", "consensus", "reason"}


def test_score_below_threshold(tmp_path):
    pytest.importorskip("numpy")
    a, b = hand_example()
    pa, pb = _pair(tmp_path, a, b)
    od = tmp_path / "o"
    assert run_tool(PACK, "score", pa, pb, "--threshold", "0.61", "--boot", "100", "--out-dir", od).returncode == 0
    assert read_summary(od)[0]["flag"] == "below threshold"


def test_score_low_prevalence_flag(tmp_path):
    pytest.importorskip("numpy")
    a, b = hand_example()  # mean ticks 4, N=10; low_prev 0.5 -> cut-off 5 -> low prevalence
    pa, pb = _pair(tmp_path, a, b)
    od = tmp_path / "o"
    assert run_tool(PACK, "score", pa, pb, "--low-prev", "0.5", "--boot", "100", "--out-dir", od).returncode == 0
    assert "low prevalence" in read_summary(od)[0]["flag"]


def test_score_constant_code_nan(tmp_path):
    pytest.importorskip("numpy")
    pa, pb = _pair(tmp_path, [0] * 10, [0] * 10)
    od = tmp_path / "o"
    r = run_tool(PACK, "score", pa, pb, "--low-prev", "0", "--boot", "100", "--out-dir", od)
    assert r.returncode == 0, r.stderr
    row = read_summary(od)[0]
    assert row["kappa"] == "NaN" and "single value" in row["note"]


def test_score_pair_and_merged(tmp_path):
    pytest.importorskip("numpy")
    a, b = hand_example()
    p1 = _pair(tmp_path, a, b, name="q1")
    sub_dir = tmp_path / "s"
    sub_dir.mkdir()
    p2 = _pair(sub_dir, a, b, name="q2")
    od = tmp_path / "o"
    r = run_tool(PACK, "score", *p1, "--pair", *p2, "--boot", "100", "--out-dir", od)
    assert r.returncode == 0, r.stderr
    rows = read_summary(od)
    assert rows[-1]["question"] == "pooled" and abs(float(rows[-1]["kappa"]) - 0.583) < 0.001
    assert rows[-1]["N"] == "20"


def test_score_pid_mismatch_exit2(tmp_path):
    pytest.importorskip("numpy")
    a, b = hand_example()
    pa, pb = _pair(tmp_path, a, b)
    write_csv(pb, ["pid", "response", "C"], [[f"Z{i}", "x", v] for i, v in enumerate(b)])
    assert run_tool(PACK, "score", pa, pb, "--boot", "10").returncode == 2


def test_bootstrap_deterministic_and_nan_safe(cp):
    import numpy as np
    c = cp.counts_matrix([1, 1, 0, 0, 1, 0], [1, 0, 0, 0, 1, 1])
    r1 = cp.bootstrap_ci(c, 200, np.random.default_rng(1))
    r2 = cp.bootstrap_ci(c, 200, np.random.default_rng(1))
    assert r1 == r2
    allzero = cp.counts_matrix([0] * 5, [0] * 5)
    lo, hi = cp.bootstrap_ci(allzero, 100, np.random.default_rng(1))
    assert lo != lo and hi != hi  # NaN, no crash


def test_point_kappa_matches_hand_value(cp):
    a, b = hand_example()
    assert abs(cp.point_kappa(a, b) - 0.58333) < 0.0001
    assert cp.point_kappa([0, 0, 0], [0, 0, 0]) != cp.point_kappa([0, 0, 0], [0, 0, 0])  # NaN


def test_selftest():
    pytest.importorskip("numpy")
    r = run_tool(PACK, "--selftest")
    assert r.returncode == 0, r.stderr + r.stdout
    assert "selftest passed" in r.stdout
