"""Tests for tools/submit/template_diff.py (needs python-docx; skipped when missing).

Each case is a difference Word's window hides but an editorial format check sees, plus the
inheritance situations that must not raise a false alarm."""
import json

import pytest

from conftest import run_tool

docx = pytest.importorskip("docx")
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402


def make(path, margin=2.54, ea_font=None, size=None, normal_ascii=None, bold=None):
    d = Document()
    s = d.sections[0]
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(margin)
    st = d.styles["Normal"]
    if normal_ascii:
        st.font.name = normal_ascii
    if ea_font:
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), ea_font)
    if size:
        st.font.size = Pt(size)
    if bold is not None:
        st.font.bold = bold
    d.add_paragraph("Body text")
    d.add_heading("Heading", level=1)
    d.save(str(path))


@pytest.fixture
def pair(tmp_path):
    return tmp_path / "m.docx", tmp_path / "t.docx"


def run_cli(pair, *args):
    return run_tool("submit/template_diff.py", pair[0], pair[1], *args)


def row(out, item):
    for line in out.splitlines():
        if line.startswith(f"| {item} "):
            return line
    pytest.fail(f"row not found: {item}\n{out}")


def test_identical_exit0(pair):
    make(pair[0]); make(pair[1])
    r = run_cli(pair)
    assert r.returncode == 0, r.stdout + r.stderr
    assert r.stdout.strip().splitlines()[-1] == "Mismatches: 0"


def test_margin_differs(pair):
    make(pair[0], margin=2.0); make(pair[1], margin=2.54)
    r = run_cli(pair)
    assert r.returncode == 1
    line = row(r.stdout, "Section 1 top margin")
    assert "no |" in line and "2 cm" in line and "2.54 cm" in line
    assert r.stdout.strip().splitlines()[-1] == "Mismatches: 4"   # top, bottom, left, right


def test_mismatches_sorted_first(pair):
    make(pair[0], margin=2.0); make(pair[1])
    r = run_cli(pair)
    rows = [l for l in r.stdout.splitlines() if l.startswith("| ") and not l.startswith("| Item")]
    flags = [l.rstrip().endswith("yes |") for l in rows]
    assert flags == sorted(flags)     # False (mismatch) rows all come before True rows


def test_eastasia_font_differs(pair):
    make(pair[0], ea_font="標楷體"); make(pair[1], ea_font="新細明體")
    r = run_cli(pair)
    assert r.returncode == 1
    line = row(r.stdout, 'Style "Normal" font (eastAsia)')
    assert "標楷體" in line and "新細明體" in line and line.endswith("no |")
    assert row(r.stdout, 'Style "Normal" font (ascii)').endswith("yes |")


def test_size_differs(pair):
    make(pair[0], size=12); make(pair[1], size=10)
    r = run_cli(pair)
    assert r.returncode == 1
    line = row(r.stdout, 'Style "Normal" font size')
    assert "12 pt" in line and "10 pt" in line and line.endswith("no |")


def test_both_unset_is_consistent(pair):
    make(pair[0]); make(pair[1])
    r = run_cli(pair)
    line = row(r.stdout, 'Style "Normal" first-line indent')
    assert "not set (Word default applies)" in line
    assert line.endswith("yes |"), line
    assert row(r.stdout, 'Style "Normal" bold').endswith("yes |")


def test_set_vs_unset_is_mismatch(pair):
    make(pair[0], bold=True); make(pair[1])
    r = run_cli(pair)
    line = row(r.stdout, 'Style "Normal" bold')
    assert "not set" in line and line.endswith("no |")


def test_inherits_through_basedon(pair):
    # Heading 1 sets no eastAsia font itself: it must resolve through docDefaults' theme font,
    # and identical sides must come out consistent
    make(pair[0]); make(pair[1])
    r = run_cli(pair)
    line = row(r.stdout, 'Style "Heading 1" font (eastAsia)')
    assert line.endswith("yes |"), line
    assert "not set" not in line


def test_style_missing_in_template(pair):
    make(pair[0]); make(pair[1])
    d = Document(str(pair[0]))
    d.add_paragraph("Quoted", style="Quote")
    d.save(str(pair[0]))
    d = Document(str(pair[1]))
    d.styles.element.remove([s for s in d.styles.element
                             if s.tag == qn("w:style") and s.find(qn("w:name")).get(qn("w:val")) == "Quote"][0])
    d.save(str(pair[1]))
    r = run_cli(pair)
    assert r.returncode == 1
    assert "Quote" in row(r.stdout, "styles used by the manuscript but missing from the template")


def test_page_field_and_json(pair):
    make(pair[0]); make(pair[1])
    r = run_cli(pair, "--json")
    j = json.loads(r.stdout)
    assert j["mismatch"] == 0
    assert any(x["item"] == "page number field (PAGE)" for x in j["rows"])


def test_read_error_exit2(pair):
    pair[0].write_bytes(b"junk"); make(pair[1])
    assert run_cli(pair).returncode == 2
