#!/usr/bin/env python3
"""source_trace.py: sentence by sentence, can a rewritten draft be traced back to its sources?

Why it exists: an AI rewrite of a co-author's manuscript added a dozen statements
the original never made, two of them method claims for procedures never carried
out. The co-author found them only by comparing the two texts sentence by
sentence. The rule is "every statement of fact, method or authorial claim must
point back to the original manuscript, the co-author's replies or the proposal;
if it cannot, do not write it or mark it for the author", but nobody reliably
compares 240 sentences by hand. This tool does the comparison mechanically:
sentences with no source are listed, and those carrying numbers or method words
are flagged red.

Each draft sentence is compared with every source sentence (best match; Chinese
by character bigrams, English by words and word pairs; the source window
includes the neighbouring sentences so split and merged sentences still match):
  traced     coverage >= 0.70   the sentence points back to a source
  rewritten  0.40 to 0.70       reworded; a human should check the meaning
  untraced   < 0.40             no source found
Flags:
  🔴 untraced and about the study itself (numbers, method words, or site facts such
     as course / teacher / student); or a number that appears in
     no source; or a rewritten sentence that smuggles in a 12+ character span
     (Chinese) absent from the source and containing a method word or a number
  🟠 untraced literature claim (carries a citation, talks about other people's
     work): the source is the cited paper, so route it to citation verification
  🟡 any other untraced sentence (mostly connectives and transitions; still
     confirm no new claim slipped in)
A citation list "(A, 2011; B, 2009)" is not split inside the parentheses.

Usage:
  source_trace.py --orig original.docx [reply.md proposal.pdf ...] --draft new.docx [--report out.md] [--gate]
  --gate   exit 1 when any red flag remains (a pre-delivery gate)
Reads .md .txt .typ .tex, .docx (needs python-docx) and .pdf (needs pdftotext).
Standard library otherwise.
"""
import argparse
import os
import re
import subprocess
import sys

METHOD_WORDS = ["問卷", "訪談", "樣本", "受試", "參與者", "信度", "效度", "顯著", "檢定", "量表", "實驗", "前測",
                "後測", "對照組", "實驗組", "編碼", "分析", "回收", "有效", "迴歸", "相關", "平均", "標準差",
                "IRB", "倫理", "同意", "招募", "觀察", "焦點團體", "工作坊",
                "survey", "interview", "participant", "sample", "reliab", "valid", "significan", "experiment",
                "歸納", "分類", "主題分析", "紮根", "三角", "面向",
                "coded", "coding", "regression", "mean", "SD", "ANOVA", "t-test", "Likert"]
NUM = re.compile(r"(?<![A-Za-z])\d+(?:\.\d+)?%?|[一二三四五六七八九十百千]+(?:位|名|人|份|次|週|個月|年)")
CITE_YEAR = re.compile(r"[（(][^（）()]*(?:19|20)\d{2}[^（）()]*[）)]")
CITE_NARR = re.compile(r"[一-鿿A-Za-z]{1,12}(?:等人?|與[一-鿿A-Za-z]{1,12})?[（(]\s*(?:19|20)\d{2}")
SELF = re.compile(r"本研究|本課程|本文|筆者|實驗組|對照組|修課|本計畫|\b(?:we|our|this study)\b", re.I)
# Words that point at the study's own setting (course, teacher, student, site). An untraced
# sentence using them is the author's facts being supplied by someone else.
STUDY = re.compile(r"課程|教師|業師|學生|長輩|學員|現場|班級|小組|上課|本校|社區|\b(?:course|class(?:room)?|teachers?|instructors?|mentors?|students?|older adults?|learners?|on[- ]site|our (?:school|university)|community)\b", re.I)
CJK = re.compile(r"[一-鿿]")


def read_any(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            sys.exit("reading .docx needs python-docx (pip install python-docx), or convert the file to .md first")
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append("　".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".pdf":
        try:
            return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout
        except FileNotFoundError:
            sys.exit("reading .pdf needs pdftotext (poppler), or convert the file to .md first")
    with open(path, encoding="utf-8") as f:
        text = f.read()
    if ext == ".typ":
        text = re.sub(r"(?m)^\s*#(set|show|import|let|include)\b.*$", "", text)
    if ext in (".md", ".txt"):
        text = re.sub(r"(?s)^---\n.*?\n---\n", "", text)
        text = re.sub(r"(?s)<!--.*?-->", "", text)
    return text


def sentences(text):
    """Return [(line number, sentence)]. Chinese splits at 。！？；, English at . ! ? followed by a
    space; headings, tables, code fences and blank lines are skipped."""
    out = []
    for ln, line in enumerate(text.split("\n"), 1):
        line = line.strip()
        if not line or line.startswith(("#", "|", "```", "![", "<")) or re.fullmatch(r"[-=*_ ]+", line):
            continue
        for s in split_line(line):
            s = s.strip()
            if len(re.sub(r"\W", "", s)) >= 6:
                out.append((ln, s))
    return out


def split_line(line):
    """Never split inside brackets: a citation list "(A, 2011; B, 2009)" cut at the semicolon
    would leave the year outside its citation, where it reads as a number no source has."""
    parts, buf, depth = [], [], 0
    for i, ch in enumerate(line):
        buf.append(ch)
        if ch in "（(「『":
            depth += 1
        elif ch in "）)」』" and depth:
            depth -= 1
        elif depth == 0 and (ch in "。！？；" or (ch in ".!?" and line[i + 1:i + 2] == " ")):
            parts.append("".join(buf))
            buf = []
    if buf:
        parts.append("".join(buf))
    return parts


def grams(s):
    s = re.sub(r"\s+", " ", s)
    if CJK.search(s):
        t = re.sub(r"[^\w一-鿿]", "", s)
        return {t[i:i + 2] for i in range(len(t) - 1)}
    w = re.findall(r"[a-z0-9]+", s.lower())
    return set(w) | {w[i] + " " + w[i + 1] for i in range(len(w) - 1)}


def numbers(s):
    return {n for n in NUM.findall(CITE_YEAR.sub("", s))}


def novel_span(s, src_grams, min_len=12):
    """The longest run in a rewritten sentence whose bigrams are all absent from the source
    (punctuation removed first, so punctuation does not break a run). Empty when shorter than
    min_len characters. Chinese only."""
    if not CJK.search(s):
        return ""
    t = re.sub(r"[^\w一-鿿]", "", s)
    best, start = (0, 0), None
    for i in range(len(t) - 1):
        if t[i:i + 2] not in src_grams:
            if start is None:
                start = i
            if i + 2 - start > best[1] - best[0]:
                best = (start, i + 2)
        else:
            start = None
    span = t[best[0]:best[1]]
    return span if len(CJK.findall(span)) >= min_len else ""


def build_index(src_sents):
    """Each source sentence plus its neighbours forms a window, so split and merged sentences match."""
    idx = []
    for i, (fname, ln, s) in enumerate(src_sents):
        win = " ".join(x[2] for x in src_sents[max(0, i - 1):i + 2] if x[0] == fname)
        idx.append((fname, ln, s, grams(win)))
    return idx


def trace(draft_sents, idx, all_src_text):
    rows = []
    src_nums = numbers(all_src_text)
    for ln, s in draft_sents:
        g = grams(s)
        best, where = 0.0, None
        best_g = set()
        for fname, sln, ss, sg in idx:
            if not g:
                break
            c = len(g & sg) / len(g)
            if c > best:
                best, where, best_g = c, (fname, sln, ss), sg
        status = "traced" if best >= 0.70 else ("rewritten" if best >= 0.40 else "untraced")
        lost_nums = sorted(n for n in numbers(s) if n not in src_nums)
        has_fact = bool(numbers(s)) or any(w.lower() in s.lower() for w in METHOD_WORDS)
        literature = bool(CITE_YEAR.search(s) or CITE_NARR.search(s)) and not SELF.search(s)
        novel = novel_span(s, best_g) if status == "rewritten" else ""
        flag = ""
        if lost_nums or status == "untraced":
            if literature:
                flag = "🟠"         # a claim about other people's work: the source is the cited paper
            elif lost_nums or has_fact or STUDY.search(s):
                flag = "🔴"         # numbers, method or site facts about the study: only the author's materials can back them
            else:
                flag = "🟡"
        elif novel and (numbers(novel) or any(w.lower() in novel.lower() for w in METHOD_WORDS)):
            flag = "🔴"             # a rewrite that carries in a method or number the source lacks
        rows.append({"line": ln, "sent": s, "cov": best, "status": status, "flag": flag,
                     "lost_nums": lost_nums, "src": where, "novel": novel})
    return rows


def report(rows, path=None):
    n = len(rows)
    cnt = {k: sum(1 for r in rows if r["status"] == k) for k in ("traced", "rewritten", "untraced")}
    red = [r for r in rows if r["flag"] == "🔴"]
    org = [r for r in rows if r["flag"] == "🟠"]
    yel = [r for r in rows if r["flag"] == "🟡"]
    L = ["# Source trace report", "",
         f"Draft {n} sentences | traced {cnt['traced']} | rewritten {cnt['rewritten']} | untraced {cnt['untraced']} | "
         f"🔴 {len(red)} | 🟠 {len(org)} | 🟡 {len(yel)}", "",
         "🔴 = a number, method or fact about the study with no trace in the original, replies or proposal: "
         "add a source, or delete it, or mark it for the author.",
         "🟠 = a new literature claim (about other people's work): the source is the cited paper; "
         "verify each citation separately, this tool does not judge it.",
         "🟡 = an untraced ordinary sentence: confirm it is only a connective or transition and adds no claim.", ""]
    for title, items in (("🔴 Must resolve", red), ("🟠 New literature claims (verify citations)", org),
                         ("🟡 Confirm no new claim", yel)):
        L.append(f"## {title} ({len(items)})")
        for r in items:
            extra = f"; numbers with no source: {', '.join(r['lost_nums'])}" if r["lost_nums"] else ""
            if r.get("novel"):
                extra += f"; span absent from the source: 「{r['novel']}」"
            near = f"; closest: {os.path.basename(r['src'][0])} L{r['src'][1]} ({r['cov']:.2f})" if r["src"] else ""
            L.append(f"- L{r['line']}: {r['sent']}{extra}{near}")
        L.append("")
    rw = [r for r in rows if r["status"] == "rewritten" and not r["flag"]]  # already-red rows are not repeated
    L.append(f"## Reworded sentences ({len(rw)}; check the meaning did not change)")
    for r in rw:
        L.append(f"- L{r['line']} ({r['cov']:.2f}): {r['sent']}  <-  {os.path.basename(r['src'][0])} L{r['src'][1]}: {r['src'][2]}")
    text = "\n".join(L) + "\n"
    if path:
        with open(path, "w", encoding="utf-8") as f:
            f.write(text)
    return text, len(red)


def main():
    ap = argparse.ArgumentParser(description=(__doc__ or "").split("\n")[0])
    ap.add_argument("--orig", nargs="+", required=True,
                    help="the original manuscript and any other traceable sources (replies, proposal, analysis reports)")
    ap.add_argument("--draft", required=True, help="the rewritten draft to check")
    ap.add_argument("--report", default="", help="write the full report to this file (default: print it)")
    ap.add_argument("--gate", action="store_true", help="exit 1 when any red flag remains")
    a = ap.parse_args()
    src_sents, all_src = [], []
    for p in a.orig:
        t = read_any(p)
        all_src.append(t)
        src_sents += [(p, ln, s) for ln, s in sentences(t)]
    rows = trace(sentences(read_any(a.draft)), build_index(src_sents), "\n".join(all_src))
    text, n_red = report(rows, a.report or None)
    if a.report:
        head = text.split("\n")[2]
        print(head + f"\nReport: {a.report}")
    else:
        print(text)
    return 1 if (a.gate and n_red) else 0


if __name__ == "__main__":
    sys.exit(main())
